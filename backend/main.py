import os
import io
import re
import json
import sqlite3
from dotenv import load_dotenv
load_dotenv()

from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel

from services import preprocessing, language_detector, ocr, vision_ai, document_analyzer

MAX_BYTES = 10 * 1024 * 1024
DB = os.path.join(os.path.dirname(__file__), "history.db")

app = FastAPI(title="AI Handwriting Intelligence")
origins = [o.strip() for o in os.getenv("ALLOWED_ORIGINS", "*").split(",")]
app.add_middleware(CORSMiddleware, allow_origins=origins, allow_methods=["*"], allow_headers=["*"])


def save_history(res: dict):
    try:
        with sqlite3.connect(DB) as c:
            c.execute("CREATE TABLE IF NOT EXISTS results(id INTEGER PRIMARY KEY, created TEXT DEFAULT CURRENT_TIMESTAMP, language TEXT, confidence REAL, data TEXT)")
            c.execute("INSERT INTO results(language, confidence, data) VALUES (?,?,?)",
                      (res["detected_language"], res["confidence"], json.dumps(res, ensure_ascii=False)))
    except Exception:
        pass


def fail(msg, code=400):
    return JSONResponse(status_code=code, content={"success": False, "error": msg})


@app.get("/")
def root():
    return {"name": "AI Handwriting Intelligence", "status": "running", "docs": "/docs"}


@app.get("/health")
def health():
    key = os.getenv("GEMINI_API_KEY", "").strip()
    return {"status": "ok", "gemini_key_configured": bool(key and key != "your_key_here")}


async def _prepare(file: UploadFile):
    data = await file.read()
    if not data:
        return None, fail("Empty file.")
    if len(data) > MAX_BYTES:
        return None, fail("Image is too large (max 10 MB).", 413)
    try:
        return preprocessing.preprocess(data), None
    except ValueError as e:
        return None, fail(str(e))
    except Exception:
        return None, fail("Could not process this image.")


@app.post("/analyze")
async def analyze(file: UploadFile = File(...)):
    """Fast stage: preprocessing + layout + language detection only."""
    pre, err = await _prepare(file)
    if err:
        return err
    layout = document_analyzer.layout_analysis(pre["binary"])
    lang = language_detector.detect(pre["original_jpeg"])
    return {"success": True, "layout": layout, "language": lang}


@app.post("/digitize")
async def digitize(file: UploadFile = File(...)):
    pre, err = await _prepare(file)
    if err:
        return err
    warnings = []
    try:
        layout = document_analyzer.layout_analysis(pre["binary"])
        lang = language_detector.detect(pre["original_jpeg"])
        if lang.get("note"):
            warnings.append(lang["note"])
        ocr_res = ocr.recognize(pre["original_jpeg"], lang.get("codes", []))
        if ocr_res["error"]:
            warnings.append("OCR engine unavailable; Vision AI is reading the image directly.")
        ai = vision_ai.verify(pre["original_jpeg"], pre["processed_png"], ocr_res["text"], lang, layout)
        if ai is None and vision_ai.LAST_ERROR["msg"]:
            warnings.append(f"AI note: {vision_ai.LAST_ERROR['msg']}")
        steps = ["Image preprocessing", "Layout analysis", "Language detection", "Handwriting recognition"]
        if ai:
            steps += ["Context understanding", "Vision verification", "Confidence analysis"]
        else:
            steps += ["Confidence analysis (OCR only)"]
        res = document_analyzer.build_result(lang, ocr_res, ai, layout, steps, warnings)
        save_history(res)
        return res
    except Exception as e:  # never crash
        return fail(f"Unexpected error: {str(e)[:150]}", 500)


class QueryIn(BaseModel):
    question: str
    document: str


@app.post("/query")
def query(q: QueryIn):
    if not q.document.strip():
        return {"success": True, "answer": "No document text to search."}
    ans = vision_ai.answer_query(q.question, q.document)
    if ans:
        return {"success": True, "answer": ans.strip()}
    # offline fallback: regex search over the extracted text
    ql = q.question.lower()
    if "phone" in ql or "number" in ql:
        hits = re.findall(r"\+?\d[\d\s\-]{7,}\d", q.document)
    elif "date" in ql:
        hits = re.findall(r"\b\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}\b", q.document)
    else:
        return {"success": True, "answer": "AI unavailable. Full extracted text:\n" + q.document}
    return {"success": True, "answer": "\n".join(hits) if hits else "Not found in the document."}


class ExportIn(BaseModel):
    format: str
    data: dict


@app.post("/export")
def export(e: ExportIn):
    text = e.data.get("corrected_text", "")
    if e.format == "txt":
        return Response(text, media_type="text/plain; charset=utf-8",
                        headers={"Content-Disposition": "attachment; filename=handwriting.txt"})
    if e.format == "json":
        return Response(json.dumps(e.data, ensure_ascii=False, indent=2), media_type="application/json",
                        headers={"Content-Disposition": "attachment; filename=handwriting.json"})
    if e.format == "docx":
        try:
            from docx import Document
            d = Document()
            d.add_heading(e.data.get("title") or "Digitized Handwriting", 1)
            d.add_paragraph(f"Language: {e.data.get('detected_language', '')}")
            for line in text.splitlines():
                d.add_paragraph(line)
            buf = io.BytesIO()
            d.save(buf)
            return Response(buf.getvalue(),
                            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            headers={"Content-Disposition": "attachment; filename=handwriting.docx"})
        except Exception:
            return fail("DOCX export failed.", 500)
    return fail("Unknown format. Use txt, json or docx.")
